"""SpreadsheetBench task interface for CSS rollouts.

This module is the *only* place CSS touches the SpreadsheetBench environment.
It exposes a narrow ``TaskEnv`` protocol so the rest of the rollout
infrastructure (``css.rollout.batch`` / ``css.rollout.selection_eval``) is
benchmark-agnostic and unit-testable without the heavyweight ``openpyxl``
dependency or the dataset on disk.

The environment is fully self-contained under ``css.envs.spreadsheetbench``:
the prompt builders (``codegen``), code executor (``executor``), cell-value
evaluator (``evaluator``), test-case discovery (``testcases``), and dataloader
(``dataloader``) are all CSS-owned modules with no external-framework imports.
The evaluation semantics follow the official SpreadsheetBench reference
(RUCKBReasoning/SpreadsheetBench).

Design notes
------------
* **ReAct agent mode (default).** ``run_one`` uses a ReAct (Reasoning + Acting)
  agent framework vendored from Trace2Skill.  The agent interacts with the
  spreadsheet through bash tool calls (Think → Act → Observe loop) instead of
  single-shot code generation.  This naturally separates reasoning from code
  execution and avoids the Qwen3 comment-degeneration problem.
* **First Law (LLM-Code division).** ``run_one`` receives only a
  target-capable ``LLMClient`` and drives the frozen task agent through
  ``target_client.complete_target_messages(...)``. It NEVER touches an
  optimizer client — the injected client *is* the backend.
* **Result-dict shape.** ``SpreadsheetBenchEnv.run_one`` produces a result
  dict with ``hard / soft / n_cases / n_pass / n_turns / fail_reason / phase /
  cases / spreadsheet_preview / target_*_prompt``. The combined skill text is
  injected into the ReAct agent's system prompt.
* **Lazy imports.** Every ``openpyxl``-touching import lives inside a method
  body so that importing this module (and therefore the whole ``css.rollout``
  stack) works without ``openpyxl`` installed. The test double ``FakeTaskEnv``
  imports nothing heavy.
* **Per-rollout isolation.** Each rollout writes under
  ``<out_dir>/predictions/<task_id>/r<rollout_index>/`` so the K repeats of a
  task never clobber each other's artifacts or ``conversation.json``.
"""
from __future__ import annotations

import hashlib
import json
import os
import traceback
from typing import Any

from css.config import CSSConfig
from css.data.rollout import TaskResult


def _skill_hash(skill_text: str) -> str:
    """Stable short hash of the skill text — the rollout-cache validity key.

    A cached rollout result may only be reused when it was produced under the
    exact same skill document; otherwise the cached trajectory belongs to a
    different policy and must not be substituted.
    """
    return hashlib.sha256((skill_text or "").encode("utf-8")).hexdigest()[:16]


# ── Trajectory hydration ────────────────────────────────────────────────────


def hydrate_trajectory(result: dict, prediction_dir: str, task_id: str) -> dict:
    """Populate ``result["conversation"]`` from a written ``conversation.json``.

    ``run_one`` does not return the conversation in its result dict; it
    *writes* it to ``<prediction_dir>/conversation.json`` (and re-writes an
    enriched copy after evaluation). CSS needs the trajectory in the result so
    ``TaskResult.from_dict`` can map it onto ``messages``.

    If the file is present and decodes to a list, it is attached under the
    ``"conversation"`` key (which ``TaskResult.from_dict`` already understands).
    Missing / unreadable files leave ``result`` untouched. Returns ``result``.
    """
    conv_path = os.path.join(prediction_dir, "conversation.json")
    if not os.path.exists(conv_path):
        return result
    try:
        with open(conv_path, encoding="utf-8") as f:
            conversation = json.load(f)
    except Exception:  # noqa: BLE001 - corrupt/partial file: leave result as-is
        return result
    if isinstance(conversation, list):
        result["conversation"] = conversation
    return result


# ── Concrete SpreadsheetBench environment ───────────────────────────────────


class SpreadsheetBenchEnv:
    """Concrete :class:`TaskEnv` backed by the self-contained SpreadsheetBench env.

    Items can be supplied three ways (checked in order):

    1. ``items`` — an explicit ``{"train": [...], "val": [...], "test": [...]}``
       mapping (or a flat list, treated as the train split). No data loading.
    2. ``split_dir`` — an existing ``train/`` ``val/`` ``test/`` split tree,
       loaded lazily via ``SpreadsheetBenchDataLoader``.
    3. ``cfg.split_dir`` / ``cfg.data_path`` — fall back to the config's
       split / dataset settings.

    All heavy imports (``openpyxl``) happen lazily inside the methods that need
    them, so constructing / importing this class is cheap and dependency-free.
    """

    def __init__(
        self,
        cfg: CSSConfig,
        *,
        data_root: str = "",
        split_dir: str = "",
        items: dict | None = None,
    ) -> None:
        self.cfg = cfg
        self.data_root = data_root or cfg.data_root
        self.split_dir = split_dir or cfg.split_dir
        # Explicitly provided items short-circuit all dataset loading.
        if items is None:
            self._items: dict[str, list[dict]] | None = None
        elif isinstance(items, dict):
            self._items = {
                "train": list(items.get("train", [])),
                "val": list(items.get("val", [])),
                "test": list(items.get("test", [])),
            }
        else:
            # A flat list is treated as the train split.
            self._items = {"train": list(items), "val": [], "test": []}
        self._loader: Any = None  # lazily built SpreadsheetBenchDataLoader

    # ── Split accessors ──────────────────────────────────────────────────

    def _ensure_loader(self) -> Any:
        """Lazily build + set up the SpreadsheetBenchDataLoader (no-op if items given)."""
        if self._loader is not None:
            return self._loader
        # Lazy import: only needed when items were not supplied directly.
        from css.envs.spreadsheetbench.dataloader import (  # noqa: PLC0415
            SpreadsheetBenchDataLoader,
        )

        loader = SpreadsheetBenchDataLoader(
            split_dir=self.split_dir,
            data_path=self.cfg.data_path,
            split_mode="split_dir" if self.split_dir else "ratio",
            split_ratio="2:1:7",
            split_seed=self.cfg.seed,
            data_root=self.data_root,
            seed=self.cfg.seed,
        )
        # SplitDataLoader.setup() resolves split_mode and loads all splits.
        loader.setup(self.cfg.to_dict())
        self._loader = loader
        return loader

    def train_items(self) -> list[dict]:
        if self._items is not None:
            return list(self._items.get("train", []))
        return list(self._ensure_loader().train_items)

    def val_items(self) -> list[dict]:
        if self._items is not None:
            return list(self._items.get("val", []))
        return list(self._ensure_loader().val_items)

    def test_items(self) -> list[dict]:
        if self._items is not None:
            return list(self._items.get("test", []))
        return list(self._ensure_loader().test_items)

    # ── Single rollout ───────────────────────────────────────────────────

    def run_one(
        self,
        item: dict,
        skill_text: str,
        target_client: "LLMClient",  # noqa: F821
        out_dir: str,
        *,
        rollout_index: int = 0,
        epoch: int = -1,
        node_id: str = "",
    ) -> TaskResult:
        """Run one (task, rollout) via a ReAct agent and score it.

        The agent uses a Think → Act → Observe loop with a bash tool to
        manipulate the spreadsheet autonomously. After the agent completes,
        the output file is evaluated against all test cases using the same
        cell-comparison logic as before.

        Returns ``TaskResult.from_dict(result)`` with ``rollout_index`` /
        ``epoch`` / ``node_id`` stamped and the hydrated conversation attached.
        """
        from css.envs.spreadsheetbench.testcases import (  # noqa: PLC0415
            _find_test_cases,
            _auto_verify_output,
        )
        from css.envs.spreadsheetbench.codegen import (  # noqa: PLC0415
            _preview_workbook,
        )
        from css.envs.spreadsheetbench.evaluator import evaluate  # noqa: PLC0415
        from css.envs.spreadsheetbench.react_adapter import (  # noqa: PLC0415
            CSSLLMClientAdapter,
            build_system_template,
            build_task_prompt,
        )
        from css.react_agent import ReActAgent, AgentConfig  # noqa: PLC0415
        from css.react_agent.bash_tool import create_bash_tool  # noqa: PLC0415
        import shutil  # noqa: PLC0415

        task_id = str(item["id"])
        instruction = item["instruction"]
        instruction_type = item.get("instruction_type", "")
        answer_position = item.get("answer_position", "")
        answer_sheet = item.get("answer_sheet", "")
        if answer_position and answer_sheet and "!" not in answer_position:
            answer_position_eval = f"{answer_sheet}!{answer_position}"
        else:
            answer_position_eval = answer_position

        itype_lower = (instruction_type or "").lower()
        if "cell" in itype_lower:
            task_type = "cell_level"
        elif "sheet" in itype_lower:
            task_type = "sheet_level"
        else:
            task_type = "other"

        sp = item.get("spreadsheet_path", f"spreadsheet/{task_id}")
        task_dir = sp if os.path.isabs(sp) else os.path.join(self.data_root, sp)

        result: dict[str, Any] = {
            "id": task_id,
            "ok": False,
            "instruction_type": instruction_type,
            "task_type": task_type,
            "task_description": instruction,
            "phase": "setup",
            "fail_reason": "",
            "llm_ok": False,
            "code_ok": False,
            "exec_ok": False,
            "n_cases": 0,
            "n_exec_pass": 0,
            "n_pass": 0,
            "soft": 0.0,
            "hard": 0,
            "n_turns": 0,
            "cases": [],
            "error": "",
            # Rollout-cache key: which skill document produced this result.
            "skill_hash": _skill_hash(skill_text),
        }

        prediction_dir = os.path.join(
            out_dir, "predictions", task_id, f"r{rollout_index}"
        )

        try:
            # Create the prediction dir UP FRONT so every unit leaves an
            # auditable artifact dir even on an early no-test-cases / setup
            # failure — otherwise such a task silently produces no dir at all.
            os.makedirs(prediction_dir, exist_ok=True)

            cases = _find_test_cases(task_dir)
            result["n_cases"] = len(cases)
            if not cases:
                result["fail_reason"] = "no-test-cases"
                return self._finalize(
                    result, prediction_dir, task_id, rollout_index, epoch, node_id
                )

            first_input = cases[0][1]
            try:
                preview_text = _preview_workbook(first_input)
            except Exception:  # noqa: BLE001
                preview_text = "(preview failed)"

            # ── Set up ReAct agent working directory ──────────────────────
            working_dir = os.path.join(prediction_dir, "workdir")
            os.makedirs(working_dir, exist_ok=True)

            work_input = os.path.join(working_dir, "input.xlsx")
            work_output = os.path.join(working_dir, "output.xlsx")
            shutil.copy(first_input, work_input)
            if os.path.exists(work_output):
                os.remove(work_output)

            spreadsheet_content = self._get_spreadsheet_content(first_input)

            # ── Build system template with skill injection ────────────────
            system_template = build_system_template(skill_text)
            task_prompt = build_task_prompt(
                instruction=instruction,
                working_dir=os.path.abspath(working_dir),
                input_file=os.path.abspath(work_input),
                output_file=os.path.abspath(work_output),
                spreadsheet_content=spreadsheet_content,
                instruction_type=instruction_type,
                answer_position=answer_position_eval,
            )

            with open(os.path.join(prediction_dir, "spreadsheet_preview.txt"), "w") as f:
                f.write(preview_text)
            with open(os.path.join(prediction_dir, "target_system_prompt.txt"), "w") as f:
                f.write(system_template)
            with open(os.path.join(prediction_dir, "target_user_prompt.txt"), "w") as f:
                f.write(task_prompt)

            result["spreadsheet_preview"] = preview_text
            result["target_system_prompt"] = system_template
            result["target_user_prompt"] = task_prompt

            # ── Create and run ReAct agent ────────────────────────────────
            max_turns = getattr(self.cfg, "max_turns", 8)
            # Per single bash command — decoupled from the whole-rollout
            # task_timeout_s. A hung command (catastrophic regex / infinite loop)
            # is killed at this bound (with its whole process tree) so it cannot
            # stall the run; legitimate spreadsheet code finishes in seconds.
            bash_timeout = int(getattr(self.cfg, "bash_timeout_s", 180))

            adapter = CSSLLMClientAdapter(target_client, max_tokens=16384, temperature=0.0)
            bash_tool = create_bash_tool(
                working_dir,
                timeout=bash_timeout,
                forbidden_dirs=[os.path.realpath(task_dir)],
            )

            config = AgentConfig(
                max_turns=max_turns,
                verbose=False,
                system_template=system_template,
            )
            agent = ReActAgent(
                client=adapter,
                tools=[bash_tool],
                config=config,
            )

            result["phase"] = "llm"
            agent_result = agent.run(task_prompt)

            result["llm_ok"] = True
            result["n_turns"] = agent_result.total_turns
            result["code_ok"] = agent_result.success

            # ── Build the COMPLETE conversation = the faithful trajectory ──
            # The recorded conversation must be the whole trajectory the agent
            # saw: the system prompt (which carries the injected skill — strategy
            # + rules — and the ReAct action protocol), the task prompt, then
            # every turn's raw assistant output (reasoning + Action) and its
            # observation. The L1 optimizer's failure analysis (Step 1b) and
            # adherence judging (Step 4) need this whole picture to assess whether
            # the agent actually followed the strategy; a thought/observation-only
            # reconstruction silently dropped the system prompt and the task, so
            # adherence was being judged blind to what the agent was even told.
            system_prompt = agent.get_system_prompt() or system_template
            conversation: list[dict] = [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": task_prompt},
            ]

            def _append_turns(steps: list, start: int = 0) -> None:
                # ``step.thought`` is the agent's raw response (its reasoning AND
                # the Action JSON), so the action is preserved verbatim.
                for step in steps[start:]:
                    if step.thought:
                        conversation.append({"role": "assistant", "content": step.thought})
                    if step.observation:
                        conversation.append({"role": "user", "content": f"Observation: {step.observation}"})

            _append_turns(agent_result.steps)

            if not agent_result.success:
                result["fail_reason"] = f"agent-failed: {agent_result.error or 'max-turns-exceeded'}"

            # If agent completed but output doesn't exist, retry with reminder
            if agent_result.success and not os.path.exists(work_output):
                remaining = max_turns - agent_result.total_turns
                if remaining > 0:
                    retry_result = agent.continue_with_message(
                        f"[System Check] The output file was NOT created at: {os.path.abspath(work_output)}\n"
                        f"Please create the output file at the exact path specified above, "
                        f"then signal ACTION: TASK_COMPLETE again."
                    )
                    result["n_turns"] = retry_result.total_turns
                    _append_turns(retry_result.steps, agent_result.total_turns)

            # ── Save conversation ─────────────────────────────────────────
            with open(os.path.join(prediction_dir, "conversation.json"), "w") as f:
                json.dump(conversation, f, ensure_ascii=False, indent=2)

            # ── Eval: copy agent output to prediction slots and evaluate ──
            output_exists = os.path.exists(work_output)

            if output_exists:
                result["phase"] = "exec"
                result["exec_ok"] = True
                result["cases"] = []
                result["n_exec_pass"] = 0
                result["n_pass"] = 0
                enrichment_parts: list[str] = []

                for no, ip, ap in cases:
                    pred_path = os.path.join(prediction_dir, f"{no}_pred.xlsx")
                    if os.path.exists(pred_path):
                        os.remove(pred_path)

                    # For multi-case tasks: re-run agent output against each
                    # test case's answer file. The agent produced one output;
                    # we copy it to each prediction slot.
                    shutil.copy(work_output, pred_path)
                    result["n_exec_pass"] += 1

                    try:
                        ev = evaluate(pred_path, ap, instruction_type, answer_position_eval)
                    except Exception as e:  # noqa: BLE001
                        ev = {"ok": False, "reason": f"eval-exception: {type(e).__name__}: {e}"}

                    if ev["ok"]:
                        result["n_pass"] += 1
                    else:
                        if not result["fail_reason"]:
                            result["fail_reason"] = f"eval-mismatch: {ev['reason'][:200]}"
                    result["cases"].append(
                        {"no": no, "stage": "eval", "ok": ev["ok"], "reason": ev.get("reason", "")}
                    )

                    if answer_position_eval:
                        verify_report = _auto_verify_output(pred_path, ap, answer_position_eval)
                        enrichment_parts.append(
                            f"## Eval Result (case {no}): {'PASS' if ev['ok'] else 'FAIL'}\n"
                            f"{ev.get('reason', '')}\n\n{verify_report}"
                        )

                if enrichment_parts:
                    enrichment_msg = "\n\n---\n\n".join(enrichment_parts)
                    conversation.append(
                        {
                            "role": "system",
                            "content": f"[POST-EXECUTION VERIFICATION]\n\n{enrichment_msg}",
                        }
                    )
                    with open(os.path.join(prediction_dir, "conversation.json"), "w") as f:
                        json.dump(conversation, f, ensure_ascii=False, indent=2)
            else:
                if not result["fail_reason"]:
                    result["fail_reason"] = "output-not-found"

            n_cases = result["n_cases"]
            n_pass = result["n_pass"]
            result["soft"] = (n_pass / n_cases) if n_cases else 0.0
            result["hard"] = 1 if (n_cases > 0 and n_pass == n_cases) else 0
            result["ok"] = bool(result["hard"])
            if result["ok"]:
                result["fail_reason"] = ""
            return self._finalize(
                result, prediction_dir, task_id, rollout_index, epoch, node_id
            )

        except Exception as e:  # noqa: BLE001
            result["fail_reason"] = f"unexpected: {type(e).__name__}: {e}"
            result["error"] = traceback.format_exc()
            return self._finalize(
                result, prediction_dir, task_id, rollout_index, epoch, node_id
            )

    @staticmethod
    def _get_spreadsheet_content(file_path: str, max_rows: int = 5) -> str:
        """Get spreadsheet content preview (same format as Trace2Skill)."""
        try:
            import openpyxl  # noqa: PLC0415
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
            lines = []
            for i, row in enumerate(ws.iter_rows(values_only=True), 1):
                if i > max_rows:
                    lines.append(f"... ({ws.max_row - max_rows} more rows)")
                    break
                row_values = [str(cell) if cell is not None else "" for cell in row]
                lines.append(str(tuple(row_values)))
            wb.close()
            return "\n".join(lines)
        except Exception as e:  # noqa: BLE001
            return f"[Could not read spreadsheet: {e}]"

    @staticmethod
    def _finalize(
        result: dict,
        prediction_dir: str,
        task_id: str,
        rollout_index: int,
        epoch: int,
        node_id: str,
    ) -> TaskResult:
        """Hydrate the trajectory, stamp provenance, persist, and build the TaskResult.

        Writes ``result.json`` (the full computed result minus the bulky
        ``conversation``, which already lives in ``conversation.json``) so a
        resumed run can load the rollout instead of re-executing it. Best-effort:
        a persistence failure never blocks returning the result.
        """
        hydrate_trajectory(result, prediction_dir, task_id)
        result["rollout_index"] = rollout_index
        result["epoch"] = epoch
        result["node_id"] = node_id
        try:
            slim = {k: v for k, v in result.items() if k != "conversation"}
            with open(os.path.join(prediction_dir, "result.json"), "w", encoding="utf-8") as f:
                json.dump(slim, f, ensure_ascii=False, indent=2, default=str)
        except Exception:  # noqa: BLE001 - persistence is best-effort
            pass
        return TaskResult.from_dict(result)

    def load_cached_result(
        self,
        item: dict,
        out_dir: str,
        *,
        rollout_index: int,
        skill_hash: str,
    ) -> "TaskResult | None":
        """Load a previously-computed rollout result for resume, or ``None``.

        Returns a reconstructed :class:`TaskResult` only when a ``result.json``
        exists for this (task, rollout_index) under ``out_dir`` AND it was
        produced under the SAME skill (``skill_hash`` match) — otherwise the
        cached trajectory belongs to a different policy and must be re-rolled.
        The conversation is re-hydrated from ``conversation.json``. Any error
        degrades to ``None`` (re-roll), so the cache is never a failure source.
        """
        task_id = str(item.get("task_id", item.get("id", "")))
        if not task_id:
            return None
        prediction_dir = os.path.join(out_dir, "predictions", task_id, f"r{rollout_index}")
        result_path = os.path.join(prediction_dir, "result.json")
        if not os.path.exists(result_path):
            return None
        try:
            with open(result_path, encoding="utf-8") as f:
                d = json.load(f)
        except Exception:  # noqa: BLE001
            return None
        if not isinstance(d, dict) or d.get("skill_hash") != skill_hash:
            return None
        hydrate_trajectory(d, prediction_dir, task_id)
        try:
            return TaskResult.from_dict(d)
        except Exception:  # noqa: BLE001
            return None


# ── Test double ──────────────────────────────────────────────────────────────


class FakeTaskEnv:
    """Deterministic :class:`TaskEnv` test double — no openpyxl / dataset.

    ``run_one`` returns a canned :class:`TaskResult` carrying a non-empty
    2-message conversation and hard/soft scores from an injected ``scorer``.

    The ``scorer`` is ``scorer(item, rollout_index) -> hard`` (an int in
    ``{0, 1}`` or a float fraction). When it returns an int the soft score
    mirrors hard; when it returns a fraction, ``hard`` is 1 iff the fraction is
    >= 1.0. With no scorer, every rollout passes.
    """

    def __init__(self, items, scorer=None) -> None:
        if isinstance(items, dict):
            self._items = {
                "train": list(items.get("train", [])),
                "val": list(items.get("val", [])),
                "test": list(items.get("test", [])),
            }
        else:
            self._items = {"train": list(items), "val": [], "test": []}
        self.scorer = scorer

    def train_items(self) -> list[dict]:
        return list(self._items.get("train", []))

    def val_items(self) -> list[dict]:
        return list(self._items.get("val", []))

    def test_items(self) -> list[dict]:
        return list(self._items.get("test", []))

    def run_one(
        self,
        item: dict,
        skill_text: str,
        target_client: "LLMClient",  # noqa: F821 - unused; double does no LLM call
        out_dir: str,
        *,
        rollout_index: int = 0,
        epoch: int = -1,
        node_id: str = "",
    ) -> TaskResult:
        task_id = str(item.get("id", item.get("task_id", "")))
        instruction = item.get("instruction", item.get("task_description", ""))

        if self.scorer is None:
            score = 1
        else:
            score = self.scorer(item, rollout_index)

        if isinstance(score, bool) or isinstance(score, int):
            hard = 1 if int(score) >= 1 else 0
            soft = float(hard)
        else:
            soft = float(score)
            hard = 1 if soft >= 1.0 else 0

        messages = [
            {
                "role": "user",
                "content": f"[FAKE TASK {task_id}] {instruction}",
            },
            {
                "role": "assistant",
                "content": (
                    f"[FAKE ROLLOUT r{rollout_index}] "
                    f"skill_len={len(skill_text)} "
                    f"-> {'PASS' if hard else 'FAIL'}"
                ),
            },
        ]

        return TaskResult.from_dict(
            {
                "id": task_id,
                "rollout_index": rollout_index,
                "hard": hard,
                "soft": soft,
                "n_cases": 1,
                "n_pass": hard,
                "fail_reason": "" if hard else "fake-fail",
                "messages": messages,
                "n_turns": 1,
                "task_type": str(item.get("instruction_type", "")) or "other",
                "task_description": instruction,
                "instruction_type": str(item.get("instruction_type", "")),
                "epoch": epoch,
                "node_id": node_id,
            }
        )
